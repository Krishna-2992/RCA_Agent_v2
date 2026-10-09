"""Reads the ServiceNow exports: incidents, service requests and change requests.

Text is kept as written, names included: if a ticket says "contact Ed", the RCA
should say so too. Only contact details (emails, phone numbers, LAN IDs) are
scrubbed, because they never help an RCA.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook


EXCEL_EPOCH = datetime(1899, 12, 30)


# ---------------------------------------------------------------------------
# Workbook access
# ---------------------------------------------------------------------------


def read_sheet(path: Path, sheet: str | None = None) -> list[dict[str, object]]:
    """Rows as dicts keyed by the header row. `sheet=None` reads the first sheet."""

    workbook = load_workbook(path, read_only=True, data_only=True)

    try:
        worksheet = workbook[sheet] if sheet else workbook.worksheets[0]
        rows = worksheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]

        records = []
        for row in rows:
            if row is None or all(cell in (None, "") for cell in row):
                continue
            records.append({
                name: ("" if cell is None else cell)
                for name, cell in zip(header, row)
                if name
            })
        return records

    finally:
        workbook.close()


def to_datetime(value: object) -> datetime | None:
    """Excel stores these exports' dates as serial day numbers."""

    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    try:
        return EXCEL_EPOCH + timedelta(days=float(value))
    except (TypeError, ValueError):
        return None


def text_of(value: object) -> str:
    text = "" if value is None else str(value)
    return text.replace("\r\n", "\n").replace("\r", "\n").strip()


# ---------------------------------------------------------------------------
# Contact details
# ---------------------------------------------------------------------------


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")

# Only numbers that are labelled as phone numbers or written in +country form.
# A bare 10-digit run is usually a control number, not a phone.
PHONE_RE = re.compile(r"(?i)((?:phone|contact#?)\s*[:=]?\s*)\+?\d[\d\s().-]{6,}\d")
INTL_PHONE_RE = re.compile(r"(?<![\w])\+\d{10,13}(?![\w])")
LAN_ID_RE = re.compile(r"(?i)(LAN/CWS\s*:\s*)\S+")


def scrub(text: str) -> str:
    """Removes emails, phone numbers and LAN IDs; leaves everything else as written."""

    if not text:
        return text
    text = EMAIL_RE.sub("[email]", text)
    text = PHONE_RE.sub(r"\1[phone]", text)
    text = INTL_PHONE_RE.sub("[phone]", text)
    return LAN_ID_RE.sub(r"\1[id]", text)


# ---------------------------------------------------------------------------
# Incidents
# ---------------------------------------------------------------------------


# Resolution notes follow one of two templates ("3.Root Cause - ..." or
# "3.3. Root Cause - ..."), or are free text. Optional step numbering may sit
# in front of the next heading, so the lookaheads allow for it.
_STEP_NUMBER = r"(?:\d+(?:\.\d+)*\.?\s*)?"

ROOT_CAUSE_RE = re.compile(
    rf"Root\s*Cause\s*[-–:]\s*(?P<value>.*?)(?=\s*/?\s*{_STEP_NUMBER}Action\s*Taken|\Z)",
    re.I | re.S,
)

ACTION_RE = re.compile(
    rf"Action\s*Taken\s*[-–:]\s*(?P<value>.*?)(?=\s*/?\s*{_STEP_NUMBER}SLA\b|\Z)",
    re.I | re.S,
)

REFERS_TO_RE = re.compile(r"refer to (INC\d+)", re.I)
PROBLEM_RE = re.compile(r"\bPRB\d{6,8}\b")


def _tidy(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip(" ./-")


def parse_resolution(notes: str) -> tuple[str | None, str | None]:
    """(root_cause, action_taken). Free-text notes return (None, notes)."""

    root = ROOT_CAUSE_RE.search(notes)
    action = ACTION_RE.search(notes)

    root_cause = _tidy(root.group("value")) if root else None
    action_taken = _tidy(action.group("value")) if action else None

    if root is None and action is None:
        action_taken = _tidy(notes) or None

    return root_cause or None, action_taken or None


def strip_prefix(short_description: str, service: str) -> str:
    """Drops the "<service> - " or " - " prefix that self-service tickets carry.

    Only known prefixes are removed: an AJS alert also contains " - " before its
    timestamp, and that must survive.
    """

    text = short_description.strip()
    for prefix in (service, "I Need Something", ""):
        marker = f"{prefix} -" if prefix else "-"
        if text.lower().startswith(marker.lower()):
            return text[len(marker):].strip()
    return text


@dataclass(frozen=True)
class Incident:
    number: str
    parent: str | None
    priority: str
    service: str
    channel: str
    opened_at: datetime | None
    resolved_at: datetime | None
    short_description: str
    description: str
    resolution_notes: str
    root_cause: str | None
    action_taken: str | None
    resolved_by: str | None = None
    refers_to: str | None = None
    problem_refs: tuple[str, ...] = ()

    @property
    def is_alert(self) -> bool:
        return self.short_description.startswith("Incident AJS")

    @property
    def symptom_text(self) -> str:
        """What a support engineer would see: summary plus description, without repeats."""

        description = self.description.strip()
        if not description:
            return self.short_description
        head = self.short_description[:40].lower()
        if head and description.lower().startswith(head):
            return description
        return f"{self.short_description}\n{description}"


def load_incidents(path: Path, sheet: str | None = None) -> list[Incident]:
    incidents = []
    for row in read_sheet(path, sheet):
        service = text_of(row.get("Service"))
        notes = scrub(text_of(row.get("Resolution notes")))
        root_cause, action_taken = parse_resolution(notes)
        refers = REFERS_TO_RE.search(notes)

        incidents.append(Incident(
            number=text_of(row.get("Number")),
            parent=text_of(row.get("Parent Incident")) or None,
            priority=text_of(row.get("Priority")),
            service=service,
            channel=text_of(row.get("Channel")),
            opened_at=to_datetime(row.get("Created")),
            resolved_at=to_datetime(row.get("Resolved")),
            short_description=scrub(strip_prefix(text_of(row.get("Short Description")), service)),
            description=scrub(text_of(row.get("Description"))),
            resolution_notes=notes,
            root_cause=root_cause,
            action_taken=action_taken,
            resolved_by=text_of(row.get("Resolved by")) or None,
            refers_to=refers.group(1).upper() if refers else None,
            problem_refs=tuple(sorted(set(PROBLEM_RE.findall(notes)))),
        ))

    return incidents


# ---------------------------------------------------------------------------
# Service requests
# ---------------------------------------------------------------------------


# Catalog-form fields that carry routing details rather than the issue.
# Matched only on "key = value" style lines, so free text that happens to start
# with "Location Management ..." is kept.
SR_BOILERPLATE_KEYS = {
    "requested for", "assignment group", "time zone", "work location",
    "business phone", "service", "configuration item", "contact type",
    "urgency", "location", "kb article followed", "user details", "lan/cws",
    "email", "contact#", "work set up",
}

SR_BOILERPLATE_PREFIXES = (
    "request for catalog item", "what is the nature", "do you know the assignment group",
    "requestor's asset", "asset for which", "customer identity verification",
)

FIELD_SPLIT_RE = re.compile(r"\s*(?:=|:|>)\s*")
DESCRIPTION_KEY_RE = re.compile(r"^Description[^=]*=\s*", re.I)
EMAIL_HEADER_RE = re.compile(r"^\s*(From|To|Cc|Sent|Subject)\s*:.*$", re.M | re.I)


def _is_boilerplate(line: str) -> bool:
    lowered = line.lower()
    if lowered.startswith(SR_BOILERPLATE_PREFIXES):
        return True
    parts = FIELD_SPLIT_RE.split(line, maxsplit=1)
    return len(parts) == 2 and parts[0].strip().lower() in SR_BOILERPLATE_KEYS


def clean_sr_description(text: str) -> str:
    kept = []
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line or _is_boilerplate(line):
            continue
        kept.append(DESCRIPTION_KEY_RE.sub("", line))
    return "\n".join(kept).strip()


def clean_work_notes(text: str) -> str:
    """Drops pasted email headers and bare acknowledgements."""

    text = EMAIL_HEADER_RE.sub("", text)
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    meaningful = [line for line in lines if line.lower().strip(".!") not in {"ack", "acknowledged"}]
    return "\n".join(meaningful)


@dataclass(frozen=True)
class ServiceRequest:
    number: str
    opened_at: datetime | None
    closed_at: datetime | None
    service: str
    short_description: str
    description: str
    work_notes: str

    @property
    def symptom_text(self) -> str:
        if self.description and self.description.lower().startswith(self.short_description[:40].lower()):
            return self.description
        return f"{self.short_description}\n{self.description}".strip()


def load_service_requests(path: Path) -> list[ServiceRequest]:
    requests = []
    for row in read_sheet(path, "SR"):
        service = text_of(row.get("Service"))
        requests.append(ServiceRequest(
            number=text_of(row.get("Number")),
            opened_at=to_datetime(row.get("Opened")),
            closed_at=to_datetime(row.get("Closed")),
            service=service,
            short_description=scrub(strip_prefix(text_of(row.get("Short Description")), service)),
            description=scrub(clean_sr_description(text_of(row.get("Description")))),
            work_notes=scrub(clean_work_notes(text_of(row.get("Work notes")))),
        ))
    return requests


# ---------------------------------------------------------------------------
# Change requests
# ---------------------------------------------------------------------------


PROGRAM_RE = re.compile(r"\bF8RH[0-9A-Z]{4}\b", re.I)


@dataclass(frozen=True)
class ChangeRequest:
    number: str
    created_at: datetime | None
    closed_at: datetime | None
    short_description: str
    change_type: str
    risk: str
    category: str
    programs: tuple[str, ...] = field(default=())


def load_changes(path: Path) -> list[ChangeRequest]:
    changes = []
    for row in read_sheet(path, "CR"):
        summary = text_of(row.get("Short Description"))
        changes.append(ChangeRequest(
            number=text_of(row.get("Number")),
            created_at=to_datetime(row.get("Created")),
            closed_at=to_datetime(row.get("Actual end date")),
            short_description=summary,
            change_type=text_of(row.get("Type")),
            risk=text_of(row.get("Risk")),
            category=f"{text_of(row.get('Category'))} / {text_of(row.get('Subcategory'))}".strip(" /"),
            programs=tuple(sorted({p.upper() for p in PROGRAM_RE.findall(summary)})),
        ))
    return changes


# ---------------------------------------------------------------------------
# Everything from the samples workbook
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class SamplesWorkbook:
    incidents: list[Incident]
    service_requests: list[ServiceRequest]
    changes: list[ChangeRequest]


def load_samples(path: Path) -> SamplesWorkbook:
    return SamplesWorkbook(
        incidents=load_incidents(path, "INC"),
        service_requests=load_service_requests(path),
        changes=load_changes(path),
    )
